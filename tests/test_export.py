import io
import json

from openpyxl import load_workbook

from chemreg_intel.demo import demo_inputs, demo_service, demo_sources
from chemreg_intel.io import export_excel, export_json


def test_excel_export_integrity():
    payload = export_excel(demo_service().screen(demo_inputs()), demo_sources())
    workbook = load_workbook(io.BytesIO(payload), read_only=True)
    assert workbook.sheetnames == [
        "Executive Summary", "Chemical Identity", "SVHC Screening",
        "REACH Restrictions", "CLP - GHS", "Evidence Log",
        "Manual Review Queue", "Source Register",
    ]


def test_json_export_preserves_source_and_status():
    parsed = json.loads(export_json(demo_service().screen([demo_inputs()[0]])))
    assert parsed[0]["svhc"]["source"]["synthetic"] is True
    assert parsed[0]["svhc"]["status"] == "CONFIRMED MATCH"

